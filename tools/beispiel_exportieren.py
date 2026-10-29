"""Liest die Planungsvorlage und schreibt data/beispiel_kw44-53.json (für die Prüfseite)."""
import sys, json, openpyxl
from openpyxl.utils import get_column_letter as L
src = sys.argv[1]; dst = sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
ws = next(s for s in wb.worksheets if 'Saisonplanung' in s.title)
DAYS = ['Mo','Di','Mi','Do','Fr','Sa','So']
hdr = next(r for r in range(1, 10) if any(c.value == 'Hauptblock' for c in ws[r]))
H = {c.value: c.column for c in ws[hdr] if c.value}
weeks = []; w = None
for r in range(hdr + 1, ws.max_row + 1):
    lab = ws.cell(r, H['Hauptblock']).value
    if lab == 'Datum':
        mon = ws.cell(r, H['Montag']).value
        w = {'kw': ws.cell(r, H['KW']).value, 'phase': ws.cell(r, H['Phase']).value, 'hint': ws.cell(r, H['Hinweis']).value,
             'days': [{'date': None, 'kalender': '', 'termine': '', 'cells': []} for _ in DAYS]}
        for d, n in enumerate(DAYS):
            v = ws.cell(r, H[n]).value; w['days'][d]['date'] = v.isoformat() if v else None
        weeks.append(w); continue
    if not w or not lab: continue
    if lab in ('Kalender', 'Termine'):
        for d, n in enumerate(DAYS): w['days'][d][lab.lower()] = ws.cell(r, H[n]).value or ''
        continue
    for d, n in enumerate(DAYS):
        v = ws.cell(r, H[n]).value
        if v: w['days'][d]['cells'].append({'block': lab, 'lines': str(v).split('\n'), 'cell': f"{L(H[n])}{r}"})
ath = [r[0].value for r in wb['Athletinnen'].iter_rows(min_row=2) if r[0].value]
json.dump({'athletes': ath, 'weeks': weeks, 'sheetName': ws.title}, open(dst, 'w'), ensure_ascii=False)
print(len(weeks), ath)
